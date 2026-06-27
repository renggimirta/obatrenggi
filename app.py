from flask import Flask, render_template, request, jsonify, send_file
from supabase import create_client, Client
import os
from datetime import datetime
import openpyxl
from io import BytesIO
from dotenv import load_dotenv

# Load environment variables from .env file
load_dotenv()

app = Flask(__name__)

# Supabase configuration
SUPABASE_URL = os.environ.get('SUPABASE_URL')
SUPABASE_KEY = os.environ.get('SUPABASE_KEY')
supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/obat', methods=['GET'])
def get_obat():
    try:
        response = supabase.table('obat').select('*').order('tanggal', desc=True).execute()
        return jsonify({'success': True, 'data': response.data})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/obat', methods=['POST'])
def add_obat():
    try:
        data = request.json
        nama_obat = data.get('nama_obat')
        qty = data.get('qty', 1)
        
        # Check if obat already exists
        existing = supabase.table('obat').select('*').eq('nama_obat', nama_obat).execute()
        
        if existing.data:
            # Update quantity if exists
            obat_id = existing.data[0]['id']
            current_qty = existing.data[0]['qty']
            new_qty = current_qty + qty
            supabase.table('obat').update({'qty': new_qty}).eq('id', obat_id).execute()
            return jsonify({'success': True, 'message': 'Quantity updated', 'data': existing.data[0]})
        else:
            # Insert new obat
            new_data = {
                'nama_obat': nama_obat,
                'qty': qty,
                'tanggal': datetime.now().isoformat()
            }
            response = supabase.table('obat').insert(new_data).execute()
            return jsonify({'success': True, 'message': 'Obat added', 'data': response.data})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/obat/<int:obat_id>', methods=['PUT'])
def update_obat(obat_id):
    try:
        data = request.json
        response = supabase.table('obat').update({
            'nama_obat': data.get('nama_obat'),
            'qty': data.get('qty')
        }).eq('id', obat_id).execute()
        return jsonify({'success': True, 'data': response.data})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/obat/<int:obat_id>', methods=['DELETE'])
def delete_obat(obat_id):
    try:
        supabase.table('obat').delete().eq('id', obat_id).execute()
        return jsonify({'success': True, 'message': 'Obat deleted'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/export', methods=['GET'])
def export_excel():
    try:
        response = supabase.table('obat').select('*').order('tanggal', desc=True).execute()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data Obat"
        
        # Headers
        ws.append(['No', 'Nama Obat', 'Qty', 'Tanggal'])
        
        # Data
        for idx, row in enumerate(response.data, 1):
            ws.append([
                idx,
                row['nama_obat'],
                row['qty'],
                row['tanggal']
            ])
        
        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"data_obat_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/import', methods=['POST'])
def import_excel():
    try:
        if 'file' not in request.files:
            return jsonify({'success': False, 'error': 'No file provided'}), 400
        
        file = request.files['file']
        wb = openpyxl.load_workbook(file)
        ws = wb.active
        
        # Skip header row
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        
        imported = 0
        for row in rows:
            if row[1]:  # nama_obat
                nama_obat = row[1]
                qty = row[2] if row[2] else 1
                
                # Check if exists
                existing = supabase.table('obat').select('*').eq('nama_obat', nama_obat).execute()
                
                if existing.data:
                    # Update qty
                    obat_id = existing.data[0]['id']
                    current_qty = existing.data[0]['qty']
                    new_qty = current_qty + qty
                    supabase.table('obat').update({'qty': new_qty}).eq('id', obat_id).execute()
                else:
                    # Insert new
                    new_data = {
                        'nama_obat': nama_obat,
                        'qty': qty,
                        'tanggal': datetime.now().isoformat()
                    }
                    supabase.table('obat').insert(new_data).execute()
                
                imported += 1
        
        return jsonify({'success': True, 'message': f'{imported} data imported'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000, ssl_context='adhoc')
